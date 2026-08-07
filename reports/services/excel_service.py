"""
Excel Service - Export data to Excel format using openpyxl.
Caldan Construction - "Building Kenya's Future, One Project at a Time."
"""

import logging
from datetime import datetime
from decimal import Decimal

from django.http import HttpResponse

logger = logging.getLogger('caldan.reports')


class ExcelService:
    """
    Service for exporting data to Excel format using openpyxl.
    Supports multiple sheets, styling, and formatting.
    """

    @classmethod
    def export(cls, data, fields, filename, field_headers=None, sheet_name='Export'):
        """
        Export data to Excel format.

        Args:
            data: List of dictionaries or objects to export
            fields: List of field names to include
            filename: Output filename (without extension)
            field_headers: Optional custom headers
            sheet_name: Name of the worksheet

        Returns:
            HttpResponse with Excel content
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.error("openpyxl not installed. Cannot export to Excel.")
            raise ImportError(
                "openpyxl is required for Excel export. "
                "Run: pip install openpyxl"
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel sheet names max 31 chars

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0B66C3", end_color="0B66C3", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Write header
        if field_headers:
            headers = field_headers
        else:
            headers = [field.replace('_', ' ').title() for field in fields]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data
        for row_idx, item in enumerate(data, 2):
            for col_idx, field in enumerate(fields, 1):
                value = cls._get_value(item, field)
                cell = ws.cell(row=row_idx, column=col_idx, value=cls._format_value(value))
                cell.border = thin_border
                if isinstance(value, (int, float, Decimal)):
                    cell.number_format = '#,##0.00'

        # Auto-column width
        for col in ws.columns:
            max_length = 0
            column_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Freeze header row
        ws.freeze_panes = 'A2'

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        wb.save(response)

        logger.info(f"Excel export generated: {filename}.xlsx with {len(data)} rows")
        return response

    @classmethod
    def export_multi_sheet(cls, sheets_data, filename):
        """
        Export multiple sheets to a single Excel file.

        Args:
            sheets_data: Dict mapping sheet_name -> (data, fields, headers)
            filename: Output filename (without extension)

        Returns:
            HttpResponse with Excel content
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.error("openpyxl not installed. Cannot export to Excel.")
            raise ImportError("openpyxl is required for Excel export.")

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0B66C3", end_color="0B66C3", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for sheet_name, (data, fields, headers) in sheets_data.items():
            ws = wb.create_sheet(title=sheet_name[:31])

            # Write header
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Write data
            for row_idx, item in enumerate(data, 2):
                for col_idx, field in enumerate(fields, 1):
                    value = cls._get_value(item, field)
                    cell = ws.cell(row=row_idx, column=col_idx, value=cls._format_value(value))
                    cell.border = thin_border
                    if isinstance(value, (int, float, Decimal)):
                        cell.number_format = '#,##0.00'

            # Auto-column width
            for col in ws.columns:
                max_length = 0
                column_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            # Freeze header row
            ws.freeze_panes = 'A2'

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        wb.save(response)

        logger.info(f"Multi-sheet Excel export generated: {filename}.xlsx")
        return response

    @classmethod
    def _get_value(cls, item, field):
        """Get a value from a data item."""
        if isinstance(item, dict):
            return item.get(field, '')
        elif hasattr(item, field):
            return getattr(item, field, '')
        else:
            return ''

    @classmethod
    def _format_value(cls, value):
        """Format a value for Excel output."""
        if value is None:
            return ''
        elif isinstance(value, datetime):
            return value
        elif isinstance(value, bool):
            return 'Yes' if value else 'No'
        elif isinstance(value, (list, dict)):
            import json
            return json.dumps(value)
        else:
            return value